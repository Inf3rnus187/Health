"""Read a table export (CSV, Excel, JSON) as rows of ``{header: text}``.

Apps export their history in every shape: comma or semicolon CSV, UTF-8
or Latin-1, an Excel sheet with a title above the headers, a JSON list.
The first row with two or more filled cells is taken as the headers.
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import datetime, time
from typing import Any

from openpyxl import load_workbook

Row = dict[str, str]
_MIN_HEADERS = 2
_MAX_DEPTH = 3


def rows_of(data: bytes, filename: str) -> list[Row]:
    """The file's rows, keyed by their column titles (none if unreadable)."""
    try:
        return _rows(data, filename)
    except (csv.Error, ValueError, OSError, KeyError, zipfile.BadZipFile):
        return []


def _rows(data: bytes, filename: str) -> list[Row]:
    """The rows of a CSV, Excel or JSON file."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _table(_sheet(data))
    text = _decode(data)
    if name.endswith(".json") or text.lstrip()[:1] in "[{":
        try:
            return _json(json.loads(text))
        except ValueError:
            pass
    return _table(_csv(text))


def _decode(data: bytes) -> str:
    """UTF-8 (with or without BOM), else Latin-1."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _csv(text: str) -> list[list[str]]:
    """CSV cells, the delimiter guessed (comma, semicolon, tab, pipe)."""
    sample = text[:4096]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        dialect.doublequote = True  # "" inside a quoted cell (a comment)
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _sheet(data: bytes) -> list[list[str]]:
    """The first sheet's cells as text (dates in ISO)."""
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book.worksheets[0]
    return [
        [_cell(value) for value in row]
        for row in sheet.iter_rows(values_only=True)
    ]


def _cell(value: Any) -> str:
    """A spreadsheet value as text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time.min:  # a date cell: no time to invent
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    return str(value)


def _table(cells: list[list[str]]) -> list[Row]:
    """Rows under the first line with two or more filled cells."""
    start = next(
        (
            i
            for i, row in enumerate(cells)
            if sum(1 for c in row if c.strip()) >= _MIN_HEADERS
        ),
        None,
    )
    if start is None:
        return []
    head = [c.strip() or f"col{i}" for i, c in enumerate(cells[start])]
    return [
        dict(zip(head, [c.strip() for c in row], strict=False))
        for row in cells[start + 1 :]
        if any(c.strip() for c in row)
    ]


def _json(value: Any, depth: int = 0) -> list[Row]:
    """A JSON list of records (or a record holding one, even deep)."""
    if isinstance(value, dict):
        lists = [v for v in value.values() if isinstance(v, list)]
        if lists:
            return _json(lists[0], depth + 1)
        inner = [v for v in value.values() if isinstance(v, dict)]
        if inner and depth < _MAX_DEPTH:  # {"data": {"items": [...]}}
            found = _json(inner[0], depth + 1)
            if len(found) > 1 or (found and found[0]):
                return found
        return [_flat(value)]
    return [_flat(item) for item in value if isinstance(item, dict)]


def _flat(record: dict[str, Any], prefix: str = "") -> Row:
    """A record's fields as text; a nested record as ``parent.child``."""
    out: Row = {}
    for key, value in record.items():
        name = f"{prefix}{key}"
        if isinstance(value, dict) and not prefix:
            out.update(_flat(value, f"{name}."))
        elif not isinstance(value, list | dict):
            out[name] = "" if value is None else str(value)
    return out
